import pymupdf
from docx import Document


def extract_text_from_pdf(file):
    """
    Extract text from a PDF resume.
    """
    text = ""

    try:
        pdf = pymupdf.open(stream=file.read(), filetype="pdf")

        for page in pdf:
            text += page.get_text()

        pdf.close()

    except Exception as e:
        raise Exception(f"Error reading PDF: {e}")

    return clean_text(text)


def extract_text_from_docx(file):
    """
    Extract text from a DOCX resume.
    """
    text = ""

    try:
        document = Document(file)

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"

        # Also read text from tables
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"

    except Exception as e:
        raise Exception(f"Error reading DOCX: {e}")

    return clean_text(text)


def clean_text(text):
    """
    Clean extracted resume text.
    """

    # Remove unnecessary spaces
    lines = [line.strip() for line in text.splitlines()]

    # Remove empty lines
    lines = [line for line in lines if line]

    # Combine everything into readable text
    cleaned_text = "\n".join(lines)

    return cleaned_text


def extract_resume_text(file):
    """
    Detect file type and extract resume text.
    """

    if file is None:
        return ""

    file_name = file.name.lower()

    if file_name.endswith(".pdf"):
        return extract_text_from_pdf(file)

    elif file_name.endswith(".docx"):
        return extract_text_from_docx(file)

    else:
        raise ValueError(
            "Unsupported file format. Please upload a PDF or DOCX resume."
        )